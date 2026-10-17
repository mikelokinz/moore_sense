import sqlite3
import json
import math
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Database source
DB_PATH = Path(r"d:\ocean\moore\backend\data\buoy_outbox.sqlite3")
con = sqlite3.connect(DB_PATH)
cur = con.cursor()
cur.execute("SELECT payload FROM outbox ORDER BY id ASC")
rows = cur.fetchall()

print(f"Extracting {len(rows)} buoy observation payloads from database...")

STATION_INFO = {
    "OMNI-AD06": {"name": "AD06", "basin": "Arabian Sea", "region": "North Arabian Sea", "lat": 18.4950, "lon": 67.4500},
    "OMNI-AD07": {"name": "AD07", "basin": "Arabian Sea", "region": "Central Arabian Sea", "lat": 14.9314, "lon": 68.9758},
    "OMNI-AD08": {"name": "AD08", "basin": "Arabian Sea", "region": "Lakshadweep Basin", "lat": 12.0681, "lon": 68.6328},
    "OMNI-AD09": {"name": "AD09", "basin": "Arabian Sea", "region": "Eight Degree Channel", "lat": 8.1831,  "lon": 73.2983},
    "OMNI-AD10": {"name": "AD10", "basin": "Arabian Sea", "region": "Lakshadweep Archipelago", "lat": 10.3217, "lon": 72.5872},
    "OMNI-BD08": {"name": "BD08", "basin": "Bay of Bengal", "region": "North Bay of Bengal", "lat": 17.8247, "lon": 89.2358},
    "OMNI-BD09": {"name": "BD09", "basin": "Bay of Bengal", "region": "North-Central BoB", "lat": 17.4997, "lon": 89.1167},
    "OMNI-BD10": {"name": "BD10", "basin": "Bay of Bengal", "region": "Central Bay of Bengal", "lat": 16.3617, "lon": 87.9903},
    "OMNI-BD11": {"name": "BD11", "basin": "Bay of Bengal", "region": "Western Bay of Bengal", "lat": 13.5250, "lon": 84.1667},
    "OMNI-BD12": {"name": "BD12", "basin": "Andaman Sea", "region": "South Andaman Sea", "lat": 10.5153, "lon": 94.0739},
    "OMNI-BD13": {"name": "BD13", "basin": "Bay of Bengal", "region": "Central-East BoB", "lat": 13.9900, "lon": 86.9969},
    "OMNI-BD14": {"name": "BD14", "basin": "Bay of Bengal", "region": "Equatorial Indian Ocean", "lat": 6.5706,  "lon": 88.2333},
}

all_records = []
latest_records = {}
temp_depth_rows = []
sal_depth_rows = []
adcp_current_rows = []

for (payload_str,) in rows:
    data = json.loads(payload_str)
    bid = data.get("buoyId")
    info = STATION_INFO.get(bid, {"name": bid, "basin": "Indian Ocean", "region": "Indian Ocean", "lat": 0.0, "lon": 0.0})
    ts = data.get("timestamp")
    
    loc = data.get("location", {})
    lat = loc.get("latitude") or info["lat"]
    lon = loc.get("longitude") or info["lon"]
    
    tel = data.get("telemetry", {})
    met = tel.get("meteorology", {})
    waves = tel.get("waves", {})
    ocean = tel.get("ocean", {})
    raw_params = data.get("rawPayload", {}).get("parameters", {})
    
    # Helper to get raw parameter value safely
    def get_raw(name):
        p = raw_params.get(name)
        if isinstance(p, dict):
            val = p.get("value")
            if val is not None and isinstance(val, (int, float)) and math.isfinite(val):
                return round(val, 3)
        return None

    # Helper to get normalized telemetry value
    def get_norm(cat, key, round_n=2):
        val = tel.get(cat, {}).get(key, {}).get("value")
        if val is not None and isinstance(val, (int, float)) and math.isfinite(val):
            return round(val, round_n)
        return None

    air_temp = get_norm("meteorology", "airTemperature", 2) or get_raw("air_temperature")
    air_press = get_norm("meteorology", "pressure", 2) or get_raw("air_pressure")
    w_speed = get_norm("meteorology", "windSpeed", 2) or get_raw("wind_speed")
    w_speed_kts = round(w_speed * 1.94384, 1) if w_speed is not None else None
    w_dir = get_norm("meteorology", "windDirection", 1) or get_raw("wind_direction")
    w_gust = get_norm("meteorology", "windGust", 2) or get_raw("wind_gust")
    humidity = get_norm("meteorology", "humidity", 1) or get_raw("humidity")
    rainfall = get_norm("meteorology", "rainfall", 2) or get_raw("rainfall")
    radiation = get_norm("meteorology", "radiation", 1) or get_raw("irradiance") or get_raw("radiationin")
    wave_h = get_norm("waves", "waveHeight", 2) or get_raw("hm0")
    sst = get_norm("ocean", "sst", 2) or get_raw("water_temperature_001m")
    salinity = get_norm("ocean", "salinity", 2) or get_raw("salinity_001m")
    
    # Currents from ADCP
    c_spd_15 = get_raw("adcp_speed_015m")
    c_dir_15 = get_raw("adcp_direction_015m")
    c_spd_50 = get_raw("adcp_speed_050m")
    c_dir_50 = get_raw("adcp_direction_050m")
    c_spd_100 = get_raw("adcp_speed_100m")
    c_dir_100 = get_raw("adcp_direction_100m")
    
    # Subsurface depths
    wt_1 = get_raw("water_temperature_001m")
    wt_10 = get_raw("water_temperature_010m")
    wt_50 = get_raw("water_temperature_050m")
    wt_100 = get_raw("water_temperature_100m")
    wt_500 = get_raw("water_temperature_500m")
    
    sal_1 = get_raw("salinity_001m")
    sal_10 = get_raw("salinity_010m")
    sal_50 = get_raw("salinity_050m")
    sal_100 = get_raw("salinity_100m")
    sal_500 = get_raw("salinity_500m")
    
    rec = {
        "buoy_id": bid,
        "name": info["name"],
        "basin": info["basin"],
        "timestamp": ts,
        "latitude": round(lat, 4),
        "longitude": round(lon, 4),
        "air_temp": air_temp,
        "air_press": air_press,
        "wind_speed": w_speed,
        "wind_speed_kts": w_speed_kts,
        "wind_dir": w_dir,
        "wind_gust": w_gust,
        "humidity": humidity,
        "rainfall": rainfall,
        "radiation": radiation,
        "wave_height": wave_h,
        "sst": sst,
        "salinity": salinity,
        "current_spd_15m": c_spd_15,
        "current_dir_15m": c_dir_15,
        "current_spd_50m": c_spd_50,
        "current_dir_50m": c_dir_50,
        "current_spd_100m": c_spd_100,
        "current_dir_100m": c_dir_100,
        "wt_1m": wt_1,
        "wt_10m": wt_10,
        "wt_50m": wt_50,
        "wt_100m": wt_100,
        "wt_500m": wt_500,
        "sal_1m": sal_1,
        "sal_10m": sal_10,
        "sal_50m": sal_50,
        "sal_100m": sal_100,
        "sal_500m": sal_500,
        "source": "INCOIS / NIOT OMNI Real-Time Portal",
    }
    all_records.append(rec)
    latest_records[bid] = rec
    
    # Subsurface depth rows
    for depth_m, wt in [(1, wt_1), (10, wt_10), (15, get_raw("water_temperature_015m")),
                        (20, get_raw("water_temperature_020m")), (30, get_raw("water_temperature_030m")),
                        (50, wt_50), (75, get_raw("water_temperature_075m")), (100, wt_100),
                        (200, get_raw("water_temperature_200m")), (500, wt_500)]:
        if wt is not None:
            temp_depth_rows.append([bid, info["name"], ts, depth_m, wt])
            
    for depth_m, sal in [(1, sal_1), (10, sal_10), (15, get_raw("salinity_015m")),
                         (20, get_raw("salinity_020m")), (30, get_raw("salinity_030m")),
                         (50, sal_50), (75, get_raw("salinity_075m")), (100, sal_100),
                         (200, get_raw("salinity_200m")), (500, sal_500)]:
        if sal is not None:
            sal_depth_rows.append([bid, info["name"], ts, depth_m, sal])

    # ADCP current rows
    for depth_m, spd, d_deg in [(15, c_spd_15, c_dir_15), (25, get_raw("adcp_speed_025m"), get_raw("adcp_direction_025m")),
                                (35, get_raw("adcp_speed_035m"), get_raw("adcp_direction_035m")), (50, c_spd_50, c_dir_50),
                                (75, get_raw("adcp_speed_075m"), get_raw("adcp_direction_075m")), (100, c_spd_100, c_dir_100)]:
        if spd is not None or d_deg is not None:
            adcp_current_rows.append([bid, info["name"], ts, depth_m, spd, d_deg])

print(f"Processed {len(all_records)} total observation records.")
print(f"Captured {len(temp_depth_rows)} temperature profile depths.")
print(f"Captured {len(sal_depth_rows)} salinity profile depths.")
print(f"Captured {len(adcp_current_rows)} ADCP current profile records.")

# --- Create Workbook ---
wb = openpyxl.Workbook()
wb.remove(wb.active)

# Styles
navy_fill = PatternFill(start_color="0A192F", end_color="0A192F", fill_type="solid")
slate_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
title_font = Font(name="Segoe UI", size=15, bold=True, color="0F172A")
sub_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
reg_font = Font(name="Segoe UI", size=10, color="1E293B")
thin_border = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1")
)

def format_sheet_header(ws, title, subtitle):
    ws.append([title])
    ws.append([subtitle])
    ws.append([])
    ws.cell(row=1, column=1).font = title_font
    ws.cell(row=2, column=1).font = sub_font

def auto_fit(ws, max_cols):
    for col in ws.iter_cols(min_col=1, max_col=max_cols):
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            v = str(cell.value or "")
            if len(v) > max_len:
                max_len = len(v)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

# ==============================================================================
# SHEET 1: ALL BUOY TELEMETRY OBSERVATIONS (EVERY ROW FROM DATABASE)
# ==============================================================================
ws1 = wb.create_sheet(title="All_Buoy_Observations")
ws1.views.sheetView[0].showGridLines = True
format_sheet_header(
    ws1,
    "INCOIS-NIOT OMNI Moored Buoy Fleet: Complete Telemetry Observation Ledger",
    f"Complete chronological record of all {len(all_records)} in-situ telemetry observation readings pulled from INCOIS OMNI buoys"
)

headers1 = [
    "Buoy ID", "Station Name", "Ocean Basin", "Observation Timestamp (UTC)",
    "Latitude (°N)", "Longitude (°E)", "Air Temp (°C)", "Pressure (hPa)",
    "Wind Speed (m/s)", "Wind Speed (knots)", "Wind Dir (°)", "Wind Gust (m/s)",
    "Humidity (%)", "Rainfall (mm)", "Radiation (W/m²)", "Wave Height Hm0 (m)",
    "SST (°C)", "Salinity (PSU)", "ADCP Current 15m (cm/s)", "Current Dir 15m (°)",
    "ADCP Current 50m (cm/s)", "Current Dir 50m (°)", "ADCP Current 100m (cm/s)", "Current Dir 100m (°)",
    "Water Temp 1m (°C)", "Water Temp 10m (°C)", "Water Temp 50m (°C)", "Water Temp 100m (°C)", "Water Temp 500m (°C)",
    "Salinity 1m (PSU)", "Salinity 10m (PSU)", "Salinity 50m (PSU)", "Salinity 100m (PSU)", "Salinity 500m (PSU)",
    "Data Source"
]

ws1.append(headers1)
h_row = 4
for c_i in range(1, len(headers1) + 1):
    c = ws1.cell(row=h_row, column=c_i)
    c.fill = navy_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Add all 2,040 observation rows
for rec in all_records:
    row = [
        rec["buoy_id"], rec["name"], rec["basin"], rec["timestamp"],
        rec["latitude"], rec["longitude"], rec["air_temp"], rec["air_press"],
        rec["wind_speed"], rec["wind_speed_kts"], rec["wind_dir"], rec["wind_gust"],
        rec["humidity"], rec["rainfall"], rec["radiation"], rec["wave_height"],
        rec["sst"], rec["salinity"], rec["current_spd_15m"], rec["current_dir_15m"],
        rec["current_spd_50m"], rec["current_dir_50m"], rec["current_spd_100m"], rec["current_dir_100m"],
        rec["wt_1m"], rec["wt_10m"], rec["wt_50m"], rec["wt_100m"], rec["wt_500m"],
        rec["sal_1m"], rec["sal_10m"], rec["sal_50m"], rec["sal_100m"], rec["sal_500m"],
        rec["source"]
    ]
    ws1.append(row)
    cur_r = ws1.max_row
    for c_i in range(1, len(row) + 1):
        cell = ws1.cell(row=cur_r, column=c_i)
        cell.font = reg_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if cur_r % 2 == 1:
            cell.fill = zebra_fill

auto_fit(ws1, len(headers1))

# ==============================================================================
# SHEET 2: LATEST READINGS SUMMARY PER BUOY
# ==============================================================================
ws2 = wb.create_sheet(title="Latest_Buoy_Readings")
ws2.views.sheetView[0].showGridLines = True
format_sheet_header(
    ws2,
    "INCOIS OMNI Fleet: Latest Surface & Atmospheric Readings by Station",
    "Latest active observation for each of the 12 Indian Ocean moored buoy stations"
)

headers2 = [
    "Buoy ID", "Station Name", "Ocean Basin", "Latest Observation (UTC)",
    "Latitude (°N)", "Longitude (°E)", "Air Temp (°C)", "Barometric Pressure (hPa)",
    "Wind Speed (m/s)", "Wind Speed (knots)", "Wind Direction (°)", "Wind Gust (m/s)",
    "Relative Humidity (%)", "Rainfall (mm)", "Solar Radiation (W/m²)", "Wave Height Hm0 (m)",
    "Sea Surface Temp (°C)", "Surface Salinity (PSU)", "Surface Current (cm/s)"
]

ws2.append(headers2)
for c_i in range(1, len(headers2) + 1):
    c = ws2.cell(row=4, column=c_i)
    c.fill = slate_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for bid in sorted(latest_records.keys()):
    rec = latest_records[bid]
    row = [
        rec["buoy_id"], rec["name"], rec["basin"], rec["timestamp"],
        rec["latitude"], rec["longitude"], rec["air_temp"], rec["air_press"],
        rec["wind_speed"], rec["wind_speed_kts"], rec["wind_dir"], rec["wind_gust"],
        rec["humidity"], rec["rainfall"], rec["radiation"], rec["wave_height"],
        rec["sst"], rec["salinity"], rec["current_spd_15m"]
    ]
    ws2.append(row)
    cur_r = ws2.max_row
    for c_i in range(1, len(row) + 1):
        cell = ws2.cell(row=cur_r, column=c_i)
        cell.font = reg_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if cur_r % 2 == 1:
            cell.fill = zebra_fill
    ws2.cell(row=cur_r, column=1).font = bold_font

auto_fit(ws2, len(headers2))

# ==============================================================================
# SHEET 3: SUBSURFACE CTD TEMPERATURE PROFILES
# ==============================================================================
ws3 = wb.create_sheet(title="Subsurface_Temperature_CTD")
ws3.views.sheetView[0].showGridLines = True
format_sheet_header(
    ws3,
    "OMNI Inductive Mooring Line: Subsurface Water Temperature Observations",
    f"Depth-resolved CTD temperature measurements ({len(temp_depth_rows)} total depth points) from 1m to 500m"
)

headers3 = ["Buoy ID", "Station Name", "Observation Timestamp (UTC)", "Sensor Depth (m)", "Measured Water Temperature (°C)"]
ws3.append(headers3)
for c_i in range(1, len(headers3) + 1):
    c = ws3.cell(row=4, column=c_i)
    c.fill = navy_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")

for row in temp_depth_rows:
    ws3.append(row)
    cur_r = ws3.max_row
    for c_i in range(1, len(row) + 1):
        cell = ws3.cell(row=cur_r, column=c_i)
        cell.font = reg_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if cur_r % 2 == 1:
            cell.fill = zebra_fill

auto_fit(ws3, len(headers3))

# ==============================================================================
# SHEET 4: SUBSURFACE CTD SALINITY PROFILES
# ==============================================================================
ws4 = wb.create_sheet(title="Subsurface_Salinity_CTD")
ws4.views.sheetView[0].showGridLines = True
format_sheet_header(
    ws4,
    "OMNI Inductive Mooring Line: Subsurface Salinity Profile Observations",
    f"Depth-resolved inductive conductivity-salinity measurements ({len(sal_depth_rows)} total depth points) from 1m to 500m"
)

headers4 = ["Buoy ID", "Station Name", "Observation Timestamp (UTC)", "Sensor Depth (m)", "Measured Salinity (PSU)"]
ws4.append(headers4)
for c_i in range(1, len(headers4) + 1):
    c = ws4.cell(row=4, column=c_i)
    c.fill = navy_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")

for row in sal_depth_rows:
    ws4.append(row)
    cur_r = ws4.max_row
    for c_i in range(1, len(row) + 1):
        cell = ws4.cell(row=cur_r, column=c_i)
        cell.font = reg_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if cur_r % 2 == 1:
            cell.fill = zebra_fill

auto_fit(ws4, len(headers4))

# ==============================================================================
# SHEET 5: ADCP CURRENT VELOCITY PROFILES
# ==============================================================================
ws5 = wb.create_sheet(title="ADCP_Current_Profiles")
ws5.views.sheetView[0].showGridLines = True
format_sheet_header(
    ws5,
    "Acoustic Doppler Current Profiler (ADCP): Multi-Depth Ocean Currents",
    f"Eulerian horizontal current speed (cm/s) and heading (°) at 15m, 25m, 35m, 50m, 75m, 100m ({len(adcp_current_rows)} records)"
)

headers5 = ["Buoy ID", "Station Name", "Observation Timestamp (UTC)", "Bin Depth (m)", "Current Speed (cm/s)", "Current Direction (°)"]
ws5.append(headers5)
for c_i in range(1, len(headers5) + 1):
    c = ws5.cell(row=4, column=c_i)
    c.fill = navy_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")

for row in adcp_current_rows:
    ws5.append(row)
    cur_r = ws5.max_row
    for c_i in range(1, len(row) + 1):
        cell = ws5.cell(row=cur_r, column=c_i)
        cell.font = reg_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if cur_r % 2 == 1:
            cell.fill = zebra_fill

auto_fit(ws5, len(headers5))

# Save workbooks
output_targets = [
    Path(r"d:\ocean\moore\buoy_data.xlsx"),
    Path(r"d:\ocean\moore\realtime_buoy_monitor\buoy_data.xlsx"),
    Path(r"d:\ocean\moore\backend\data\buoy_data.xlsx"),
]

for target in output_targets:
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    print(f"Generated: {target} ({target.stat().st_size:,} bytes)")
